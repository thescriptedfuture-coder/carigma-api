"""Resume and LinkedIn-export extraction. Ported from V1's `resumeiq/parsing.py`.

The parsing is the IP, so the prompts are carried across **verbatim**. Same
libraries too — a different PDF reader extracts subtly different text, and the
analyst downstream would see a different profile for the same file.

## What V1 does, and what this keeps

Three inputs, one output shape:

    a PDF   ─┐
    a .docx ─┼─→ raw text ─→ `parse_profile_text` ─→ the profile fields
    pasted  ─┘

The paste path exists because a scanned resume produces no text and a LinkedIn
export is not something everyone can produce on a phone. It is not a fallback
for failure — it is a first-class input, and V1 learned that the hard way.

## What is deliberately NOT ported

`fetch_linkedin_profile` — the Proxycurl URL sync. It was removed from V1's UI
(bugfix #1: "PDF only") because it needs a paid per-lookup key that is not
configured, and a control that cannot work is the §25 bug class. Porting dead
code so V2 can also not use it would be worse than leaving it in V1's history.

## The size limit is here, not only in the route

A 40MB PDF is a memory problem before it is a validation problem, and the
service is the layer both the route and any future worker share.
"""

from __future__ import annotations

import io
import logging
from typing import Any, Protocol

logger = logging.getLogger(__name__)

#: Bigger than any real resume and small enough that ten at once is survivable.
MAX_BYTES = 10 * 1024 * 1024

#: V1's threshold, kept. Below this there is not enough to extract anything
#: honest from, and guessing at a near-empty document is how a profile gets
#: invented.
MIN_TEXT = 80

#: What the extractor is given. V1 truncates at the same point.
MAX_PROMPT_CHARS = 12000

#: The fields the extractor returns. Named here so the route can check the
#: response against a declaration rather than trusting the model's key
#: spelling — the renaming-boundary rule, one boundary further out.
EXTRACTED_FIELDS: tuple[str, ...] = (
    "name",
    "currentRole",
    "linkedinHeadline",
    "aboutSection",
    "skills",
    "experience",
    "education",
    "certifications",
    "industry",
    "targetRoles",
    "location",
)


class Unreadable(ValueError):
    """The file or text could not be turned into enough words to work with.

    Carries a sentence the user can act on. Not an error in the "something
    broke" sense — a scanned PDF is a normal thing to upload, and the answer
    is "paste it instead", not an apology.
    """


class NotAProfile(ValueError):
    """Readable, and not a description of anybody's working life.

    An invoice, a ticket, a bank statement. Reading one and producing a career
    score from it would be the product confidently describing something it
    never saw — the honesty rule applied to what comes IN, not only to what
    goes out. Distinct from `Unreadable` because the advice is different: the
    file was fine, it was the wrong file.
    """


class ProfileCheckFailed(RuntimeError):
    """The check itself gave no verdict. OUR failure, never the user's file.

    Deliberately not a `NotAProfile`: telling someone their real resume is not
    a resume because a model call came back malformed would blame them for us.
    And deliberately not a pass: a gate that opens when it cannot decide is not
    a gate.
    """


class Extractor(Protocol):
    """The JSON-returning model call. Injected so tests never hit the API."""

    def __call__(self, system: str, user: str, max_tokens: int) -> dict[str, Any]: ...


def extract_pdf_text(data: bytes) -> str:
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
    return "\n\n".join(parts).strip()


def extract_docx_text(data: bytes) -> str:
    """Paragraphs AND tables — V1 learned that resumes hide half their content
    in table cells, and a paragraphs-only read loses entire employment
    histories without appearing to fail."""
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(x for x in parts if x and x.strip()).strip()


def text_from_upload(data: bytes, filename: str) -> str:
    """Route a file to the right reader by extension, as V1 does.

    A failure here is `Unreadable` with the sentence for that input type, never
    a raw library exception: "PdfminerException" tells the user nothing they
    can act on, and the useful advice differs by format.
    """
    if len(data) > MAX_BYTES:
        raise Unreadable(
            "That file is larger than 10MB. A resume export is usually well under 1MB — "
            "if yours is scanned images, paste your profile text instead."
        )
    name = (filename or "").lower()
    try:
        text = extract_docx_text(data) if name.endswith(".docx") else extract_pdf_text(data)
    except Unreadable:
        raise
    except Exception as exc:
        logger.info("extraction failed for %s: %s", name or "(unnamed)", exc)
        raise Unreadable(
            "We couldn't read that file. Try re-exporting from LinkedIn "
            "(Profile → More → Save to PDF), or paste your profile text instead."
        ) from exc
    return text


#: How much of the document the profile check reads. A resume says what it is
#: in its first screen; an invoice does too.
CHECK_CHARS = 4000

PROFILE_CHECK_SYSTEM = (
    "You decide whether a document describes one person's working life. "
    "Return ONLY a raw JSON object. No prose, no markdown. Start with { end with }."
)


def check_is_profile(text: str, *, extractor: Extractor, pasted: bool) -> None:
    """Refuse text that is not a work profile, BEFORE anything is extracted.

    A separate call rather than a key added to the extraction prompt, for two
    reasons. The extraction prompt is V1's, carried verbatim because the
    parsing is the IP. And it is written to be generous — "infer one if not
    explicit" — which is right for a real profile and exactly wrong for an
    invoice: asked to infer a headline, a model will find one.

    **Biased toward yes.** Turning away someone's real resume is a closed door
    at the first step; reading an unusual one costs nothing. So short, junior,
    non-English and oddly formatted profiles pass, and only a document that is
    clearly something else is refused.

    Raises `NotAProfile` with a sentence for the input the person used, or
    `ProfileCheckFailed` when the answer is not a verdict.
    """
    user = f"""Does the text below describe a person's professional profile: a resume, CV, LinkedIn profile export, professional bio or portfolio summary covering their work, roles, skills or education?

Count it as a profile however short, junior, unusual, badly formatted or non-English it is. A student with no jobs yet, a career changer and a one-paragraph bio all count.

It is NOT a profile when it is clearly another kind of document: an invoice, receipt, ticket, bank or card statement, contract, letter, form, article, report, job advertisement, or a document about a company rather than a person.

When unsure, answer true. Wrongly turning away someone's real resume costs far more than reading an unusual one.

Return exactly: {{"is_work_profile": true or false, "looks_like": "two to four words naming the kind of document"}}

DOCUMENT TEXT:
{text[:CHECK_CHARS]}"""

    # `object`, not the protocol's dict: the model can return a list, and this
    # is the line that has to notice.
    raw: object = extractor(PROFILE_CHECK_SYSTEM, user, 200)
    verdict = raw.get("is_work_profile") if isinstance(raw, dict) else None

    # `is True` / `is False`, not truthiness. "false" as a string is truthy, and
    # a check that reads it as yes has just waved an invoice through.
    if verdict is True:
        return
    if verdict is False:
        looks_like = str(raw.get("looks_like") or "")[:60] if isinstance(raw, dict) else ""
        # What it resembled, never the document itself: this is someone's file.
        logger.info("refused a non-profile upload (looks like: %s)", looks_like or "unstated")
        if pasted:
            raise NotAProfile(
                "This doesn't look like a work profile — paste the text of your LinkedIn "
                "profile or resume, or upload your LinkedIn PDF export."
            )
        raise NotAProfile(
            "This doesn't look like a work profile — try your LinkedIn PDF export "
            "(Profile → More → Save to PDF) or a resume."
        )
    raise ProfileCheckFailed(f"the profile check returned no verdict: {str(raw)[:200]}")


def parse_profile_text(text: str, *, extractor: Extractor, pasted: bool = False) -> dict[str, Any]:
    """Raw text to structured fields. V1's prompt, unchanged.

    The `targetRoles` and `location` keys are in the resume prompt but not in
    V1's LinkedIn-PDF prompt; this uses the resume one for every input, which
    is what V1's `parse_resume` and its paste path both do.

    Three refusals come first, cheapest first: too little text, then not a
    profile, and only then the extraction.
    """
    text = (text or "").strip()
    if len(text) < MIN_TEXT:
        raise Unreadable(
            "Couldn't read enough text. Use a text-based PDF or Word resume "
            "(a scanned image won't work), or paste more of your profile."
        )

    check_is_profile(text, extractor=extractor, pasted=pasted)

    system = (
        "You extract structured career data from a resume or LinkedIn export. "
        "Return ONLY a raw JSON object. No prose, no markdown. Start with { end with }."
    )
    user = f"""Extract profile data from this resume / LinkedIn text. Return a JSON object with EXACT keys:
- name (string)
- currentRole (string — latest job title AND company)
- linkedinHeadline (string — a professional tagline; infer one if not explicit)
- aboutSection (string — a summary/objective if present, else "")
- skills (string — comma-separated list of ALL skills/tools mentioned)
- experience (string — detailed summary of work history with achievements)
- education (string — degrees and institutions)
- certifications (string — licenses, certs, courses with issuers/dates)
- industry (string — inferred from their roles)
- targetRoles (string — roles they appear to be targeting; infer from recent titles if unclear)
- location (string — city/country if present, else "")

If a field is missing, use empty string "". Return ONLY the JSON object.

RESUME / PROFILE TEXT:
{text[:MAX_PROMPT_CHARS]}"""

    raw = extractor(system, user, 3000)
    return clean(raw)


def clean(raw: dict[str, Any] | None) -> dict[str, Any]:
    """Keep the fields we declared, as strings, dropping the empty ones.

    Two rules, both learned elsewhere in this codebase:

    **Unknown keys are dropped, not merged.** The model is asked for an exact
    set; anything else is a hallucinated field, and `profile_to_db` would
    (correctly) raise on it. Dropping here means a stray key costs a field, not
    the whole upload.

    **Empty strings are dropped rather than written.** The prompt asks for `""`
    on a missing field, and writing that over an existing value would let a
    thin re-upload erase a profile the user had already filled in.
    """
    out: dict[str, Any] = {}
    for key in EXTRACTED_FIELDS:
        value = (raw or {}).get(key)
        if isinstance(value, str) and value.strip():
            out[key] = value.strip()
    return out


__all__ = [
    "CHECK_CHARS",
    "EXTRACTED_FIELDS",
    "MAX_BYTES",
    "MIN_TEXT",
    "PROFILE_CHECK_SYSTEM",
    "Extractor",
    "NotAProfile",
    "ProfileCheckFailed",
    "Unreadable",
    "check_is_profile",
    "clean",
    "extract_docx_text",
    "extract_pdf_text",
    "parse_profile_text",
    "text_from_upload",
]
